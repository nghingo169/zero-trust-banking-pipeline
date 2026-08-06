from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
SOURCE = ROOT / "source.md"
OUT = ROOT.parent / "Synthetic_Banking_Mock_Data_Design_and_Delivery_Guide.docx"
FIRST_WORKFLOW_IMAGE = Path("/var/folders/y1/5s9yw3vx22dc7v8kc1cr_dxr0000gn/T/codex-clipboard-c26ec2bd-a81d-4f7b-bb8a-ad850fa2aaff.png")
RELATIONSHIP_MODEL_IMAGE = Path("/var/folders/y1/5s9yw3vx22dc7v8kc1cr_dxr0000gn/T/codex-clipboard-73c69eef-d3b3-4d0f-846a-a382387d0bc5.png")

INK = "172B4D"
BLUE = "1F4E79"
LIGHT_BLUE = "DCE6F1"
GRAY = "F3F5F7"
CODE_BG = "F2F4F7"

def set_font(run, name="Aptos", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar"); tcPr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}"); mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); tag = OxmlElement("w:tblHeader"); tag.set(qn("w:val"), "true"); trPr.append(tag)

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:cantSplit")
    trPr.append(tag)

def set_table_geometry(table, widths_inches):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    total = int(sum(widths_inches) * 1440)
    tblW.set(qn("w:w"), str(total)); tblW.set(qn("w:type"), "dxa")
    layout = tblPr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout"); tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for gc, width in zip(grid.gridCol_lst, widths_inches): gc.set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_inches):
            cell.width = Inches(width)
            tcW = cell._tc.tcPr.tcW
            tcW.set(qn("w:w"), str(int(width * 1440))); tcW.set(qn("w:type"), "dxa")

def clean(text):
    text = text.replace("<br>", "\n").replace("<br/>", "\n")
    return text.replace("\\", "").strip()

def add_inline(p, text, size=10.5):
    # modest Markdown inline support: bold, code, and links preserve readable text
    pattern = r"(\*\*.*?\*\*|`[^`]+`|\[[^]]+\]\([^)]*\))"
    chunks = re.split(pattern, text)
    for chunk in chunks:
        if not chunk: continue
        if chunk.startswith("**") and chunk.endswith("**"):
            r = p.add_run(chunk[2:-2]); set_font(r, size=size, bold=True)
        elif chunk.startswith("`") and chunk.endswith("`"):
            r = p.add_run(chunk[1:-1]); set_font(r, name="Consolas", size=max(8, size-0.5), color=INK)
        elif chunk.startswith("["):
            m = re.match(r"\[([^]]+)\]\(([^)]*)\)", chunk)
            r = p.add_run(m.group(1) if m else chunk); set_font(r, size=size, color=BLUE)
            r.underline = True
        else:
            r = p.add_run(chunk); set_font(r, size=size)

def parse_table(lines):
    rows=[]
    for line in lines:
        if re.match(r"^\|\s*[-:]+", line): continue
        rows.append([clean(x) for x in line.strip().strip("|").split("|")])
    return rows

def table_widths(headers, count):
    if count == 8:
        return [0.85, 0.70, 0.65, 0.90, 1.10, 0.90, 2.35, 1.92]
    if count == 5: return [1.65, 1.55, 1.55, 2.25, 2.37]
    if count == 4: return [1.7, 2.0, 2.7, 2.97]
    if count == 3: return [2.25, 3.85, 3.27]
    if count == 2: return [2.4, 6.97]
    return [9.37 / count] * count

def add_table(doc, rows):
    cols = len(rows[0]); table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    widths = table_widths(rows[0], cols)
    for c, value in zip(table.rows[0].cells, rows[0]):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; shade(c, LIGHT_BLUE); set_cell_margins(c)
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, value, size=8); [setattr(r, 'bold', True) for r in p.runs]
    set_repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for data in rows[1:]:
        cells = table.add_row().cells
        for idx, (c, value) in enumerate(zip(cells, data)):
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(c)
            if idx == 0: shade(c, GRAY)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
            add_inline(p, value, size=7.6 if cols >= 8 else 8.5)
        prevent_row_split(table.rows[-1])
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_code(doc, language, lines):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(0)
    r=p.add_run((language or "text").upper()); set_font(r, size=8, bold=True, color="5B6573")
    table=doc.add_table(rows=1, cols=1); table.alignment=WD_TABLE_ALIGNMENT.LEFT; table.style="Table Grid"
    cell=table.cell(0,0); shade(cell, CODE_BG); set_cell_margins(cell, 100, 120, 100, 120)
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run("\n".join(lines)); set_font(r, name="Consolas", size=7.5, color="293241")
    set_table_geometry(table, [9.37])
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def main():
    doc=Document()
    sec=doc.sections[0]
    sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Inches(11); sec.page_height=Inches(8.5)
    sec.top_margin=Inches(.6); sec.bottom_margin=Inches(.6); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
    sec.header_distance=Inches(.3); sec.footer_distance=Inches(.3)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Aptos'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.15
    for name, size, color, before, after in [('Heading 1',16,BLUE,14,7),('Heading 2',13,BLUE,12,6),('Heading 3',11,INK,10,5)]:
        s=styles[name]; s.font.name='Aptos Display'; s._element.rPr.rFonts.set(qn('w:ascii'),'Aptos Display'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=True; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
    # Header/footer
    hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=hp.add_run('SYNTHETIC BANKING • DESIGN & DELIVERY GUIDE'); set_font(r,size=8,bold=True,color='65758B')
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=fp.add_run('Confluence-ready reference • Synthetic data only'); set_font(r,size=8,color='65758B')
    lines=SOURCE.read_text().splitlines(); i=0; first_title=True; mermaid_count=0
    while i < len(lines):
        line=lines[i]
        if line.startswith('```'):
            lang=line[3:].strip(); i+=1; block=[]
            while i<len(lines) and not lines[i].startswith('```'):
                block.append(lines[i]); i+=1
            if lang == 'mermaid':
                mermaid_count += 1
                if mermaid_count in (1, 2, 4):
                    image_path = RELATIONSHIP_MODEL_IMAGE if mermaid_count == 2 else FIRST_WORKFLOW_IMAGE
                    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
                    p.add_run().add_picture(str(image_path), width=Inches(9.37))
                else:
                    add_code(doc, lang, block)
            else:
                add_code(doc, lang, block)
            i+=1; continue
        if re.match(r'^\|', line):
            block=[]
            while i<len(lines) and lines[i].startswith('|'):
                block.append(lines[i]); i+=1
            add_table(doc, parse_table(block)); continue
        h=re.match(r'^(#{1,4})\s+(.*)', line)
        if h:
            level=len(h.group(1)); text=h.group(2)
            if level==1 and first_title:
                p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
                add_inline(p, 'Synthetic Banking Mock Data Design and Delivery Guide', size=23)
                for r in p.runs: set_font(r,name='Aptos Display',size=23,bold=True,color=INK)
                sub=doc.add_paragraph(); sub.paragraph_format.space_after=Pt(12); r=sub.add_run('Confluence-ready reference for the synthetic Bronze/source environment'); set_font(r,size=11,italic=True,color='65758B')
                first_title=False
            else:
                p=doc.add_paragraph(style=f'Heading {min(level,3)}'); add_inline(p, text, size={2:16,3:13,4:11}[level])
            i+=1; continue
        if re.match(r'^[-*] ', line):
            p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); add_inline(p, line[2:], size=10.5); i+=1; continue
        if re.match(r'^\d+\. ', line):
            p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(3); add_inline(p, re.sub(r'^\d+\.\s*','',line),size=10.5); i+=1; continue
        if line.startswith('> '):
            table=doc.add_table(rows=1,cols=1); table.style='Table Grid'; cell=table.cell(0,0); shade(cell,'FFF4CE'); set_cell_margins(cell,100,140,100,140); p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); add_inline(p,line[2:],size=10); i+=1; continue
        if line.strip():
            p=doc.add_paragraph(); add_inline(p, clean(line),size=10.5)
        i+=1
    doc.core_properties.title='Synthetic Banking Mock Data Design and Delivery Guide'
    doc.core_properties.subject='Confluence-ready synthetic banking data reference'
    doc.core_properties.author='Synthetic Banking Mock Data Generator'
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__': main()
